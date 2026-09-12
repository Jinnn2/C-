"""Q1 deterministic MILP, LP lower bound, no-storage control and exports."""
from __future__ import annotations
import csv
import hashlib
import json
import math
import sqlite3
import time
from pathlib import Path

import numpy as np
import openpyxl
import scipy
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix
from validate_q1 import validate

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'results'/'q1'
N=144
Q,C,D,W,S,Z=0,144,288,432,576,721
SIZE=865


def solve(data, relaxed=False):
    objective=np.zeros(SIZE)
    objective[Q:Q+N]=[r['price_yuan_per_kwh'] for r in data]
    lower=np.zeros(SIZE); upper=np.full(SIZE,np.inf)
    upper[C:C+N]=upper[D:D+N]=5000/6
    lower[S:S+N+1]=1200; upper[S:S+N+1]=10800
    lower[S]=upper[S]=lower[S+N]=upper[S+N]=6000
    upper[Z:]=1
    integrality=np.zeros(SIZE,dtype=int)
    if not relaxed: integrality[Z:]=1
    matrix=lil_matrix((4*N,SIZE)); lo=np.full(4*N,-np.inf); hi=np.zeros(4*N)
    for t,r in enumerate(data):
        matrix[t,Q+t]=1; matrix[t,D+t]=1; matrix[t,C+t]=-1; matrix[t,W+t]=-1
        lo[t]=hi[t]=r['load_kwh']-r['pv_forecast_kwh']
        k=N+t
        matrix[k,S+t+1]=1; matrix[k,S+t]=-1; matrix[k,C+t]=-.9; matrix[k,D+t]=1/.9
        lo[k]=hi[k]=0
        matrix[2*N+t,C+t]=1; matrix[2*N+t,Z+t]=-5000/6
        matrix[3*N+t,D+t]=1; matrix[3*N+t,Z+t]=5000/6
        hi[3*N+t]=5000/6
    started=time.perf_counter()
    res=milp(objective,integrality=integrality,bounds=Bounds(lower,upper),
             constraints=LinearConstraint(matrix.tocsr(),lo,hi),
             options={'mip_rel_gap':1e-9,'time_limit':120})
    elapsed=time.perf_counter()-started
    if not res.success: raise RuntimeError(f'Solve failed: {res.message}')
    rows=[]
    for t,r in enumerate(data):
        rows.append(dict(slot=t+1,interval_label=r['interval_label'],
                         price_yuan_per_kwh=r['price_yuan_per_kwh'],load_kwh=r['load_kwh'],
                         pv_forecast_kwh=r['pv_forecast_kwh'],purchase_kwh=float(res.x[Q+t]),
                         charge_kwh=float(res.x[C+t]),discharge_kwh=float(res.x[D+t]),
                         curtailment_kwh=float(res.x[W+t]),soc_start_kwh=float(res.x[S+t]),
                         soc_end_kwh=float(res.x[S+t+1]),cost_yuan=float(res.x[Q+t]*r['price_yuan_per_kwh'])))
    info=dict(status=int(res.status),message=str(res.message),objective_yuan=float(res.fun),
              elapsed_seconds=elapsed,relative_gap=float(getattr(res,'mip_gap',0) or 0),
              dual_bound_yuan=float(res.fun if getattr(res,'mip_dual_bound',None) is None else res.mip_dual_bound),
              node_count=int(getattr(res,'mip_node_count',0) or 0))
    return rows,info


def write_csv(path,rows):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    quality=json.loads((ROOT/'reports/data_quality_report.json').read_text(encoding='utf-8'))
    assert quality['status']=='PASS_WITH_WARNINGS', 'Data preparation must pass'
    # Fail if original files changed after preparation.
    for f in quality['sources']:
        assert hashlib.sha256((ROOT/f['path']).read_bytes()).hexdigest()==f['sha256'], f['path']
    with sqlite3.connect(f'file:{(ROOT/"data/processed/microgrid.sqlite").as_posix()}?mode=ro',uri=True) as conn:
        conn.row_factory=sqlite3.Row
        data=[dict(r) for r in conn.execute('SELECT * FROM baseline ORDER BY slot')]
    assert len(data)==144
    strict,strict_info=solve(data)
    relaxed,lp_info=solve(data,True)
    baseline=[]
    for r in strict:
        b=dict(r)
        b.update(purchase_kwh=max(r['load_kwh']-r['pv_forecast_kwh'],0),charge_kwh=0.,discharge_kwh=0.,
                 curtailment_kwh=max(r['pv_forecast_kwh']-r['load_kwh'],0),soc_start_kwh=6000.,soc_end_kwh=6000.)
        b['cost_yuan']=b['purchase_kwh']*b['price_yuan_per_kwh']; baseline.append(b)
    audits={
        'strict':validate(strict,strict_info['objective_yuan']),
        'lp_relaxation':validate(relaxed,lp_info['objective_yuan'],require_exclusive=False),
        'no_storage':validate(baseline),
    }
    assert lp_info['objective_yuan']<=strict_info['objective_yuan']+1e-5
    assert strict_info['objective_yuan']<=audits['no_storage']['cost_yuan']+1e-5
    # Check the auditor can reject a real supply defect, not just accept solver outputs.
    broken=[dict(r) for r in strict]; broken[0]['purchase_kwh']+=1
    rejected=False
    try: validate(broken)
    except AssertionError: rejected=True
    assert rejected, 'Auditor failed to reject injected imbalance'
    blocks=[]
    for a in range(0,N,24):
        rs=strict[a:a+24]
        blocks.append(dict(interval_label=f'{a//6:02d}:00-{(a+24)//6:02d}:00',
                           charge_kwh=math.fsum(r['charge_kwh'] for r in rs),
                           discharge_kwh=math.fsum(r['discharge_kwh'] for r in rs)))
    selected=[strict[h*6] for h in (10,12,14,16,18,20)]
    for name,rows in [('dispatch',strict),('lp_dispatch',relaxed),('no_storage_dispatch',baseline),
                      ('four_hour_summary',blocks),('selected_intervals',selected)]:
        write_csv(OUT/f'{name}.csv',rows)
    template=ROOT/'附件/附件5/result1.xlsx'
    workbook=openpyxl.load_workbook(template)
    s=workbook['计划购电量']
    for i,r in enumerate(strict,2):
        s.cell(i,1,r['interval_label']); s.cell(i,2,r['purchase_kwh'])
    s=workbook['充放电量']
    for i,r in enumerate(blocks,2):
        s.cell(i,1,r['interval_label']);s.cell(i,2,r['charge_kwh']);s.cell(i,3,r['discharge_kwh'])
    s['D2']='0:00';s['E2']=strict[0]['soc_start_kwh'];s['D3']='24:00';s['E3']=strict[-1]['soc_end_kwh']
    for sheet in workbook:
        sheet.freeze_panes='B2'
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,(int,float)): cell.number_format='0.000000'
        for col in ('A','B','C','D','E'): sheet.column_dimensions[col].width=22
    result_path=ROOT/'results/result1.xlsx'
    workbook.save(result_path);workbook.close()
    # Read back actual persisted artifacts, including original data columns in the CSV.
    with (OUT/'dispatch.csv').open(encoding='utf-8-sig',newline='') as f:
        csvrows=[{k:(v if k=='interval_label' else int(v) if k=='slot' else float(v)) for k,v in r.items()} for r in csv.DictReader(f)]
    validate(csvrows,strict_info['objective_yuan'])
    w=openpyxl.load_workbook(result_path,data_only=True)
    saved=list(w['计划购电量'].values)[1:]
    assert len(saved)==144 and all(a==r['interval_label'] and abs(b-r['purchase_kwh'])<1e-8 for (a,b),r in zip(saved,strict))
    for i,b in enumerate(blocks,2):
        assert abs(w['充放电量'].cell(i,2).value-b['charge_kwh'])<1e-8
        assert abs(w['充放电量'].cell(i,3).value-b['discharge_kwh'])<1e-8
    assert w['充放电量']['E2'].value==6000 and w['充放电量']['E3'].value==6000
    w.close()
    for f in quality['sources']:
        assert hashlib.sha256((ROOT/f['path']).read_bytes()).hexdigest()==f['sha256']
    savings=audits['no_storage']['cost_yuan']-audits['strict']['cost_yuan']
    summary=dict(model='144-step grid-side energy MILP',scipy_version=scipy.__version__,
                 soc_boundary_kwh=6000,charge_efficiency=.9,discharge_efficiency=.9,
                 strict_solver=strict_info,lp_solver=lp_info,audits=audits,
                 savings_yuan=savings,savings_percent=savings/audits['no_storage']['cost_yuan']*100,
                 lp_gap_yuan=strict_info['objective_yuan']-lp_info['objective_yuan'],
                 export_readback_passed=True,injected_defect_rejected=True,originals_unchanged=True,
                 inputs=quality['sources'])
    (OUT/'experiment.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    lines=['# 第一问实验报告','',
           '采用144段确定性MILP；日初日末6000 kWh，充放电各90%，电网侧5000 kW，禁止同时充放电。',
           '仅以附件1作为负载、光伏与电价输入。沿用右端点代表此前10分钟区间的假设；未做二级目标优化。','',
           '## 对照结果','', '| 策略 | 购电量/kWh | 购电费/元 | 弃用/kWh |','|---|---:|---:|---:|']
    for key,title in [('no_storage','无储能'),('lp_relaxation','线性松弛下界'),('strict','严格互斥MILP')]:
        a=audits[key];lines.append(f"| {title} | {a['purchase_kwh']:.6f} | {a['cost_yuan']:.6f} | {a['curtailed_kwh']:.6f} |")
    lines += ['',f'相对无储能节费{savings:.6f}元（{summary["savings_percent"]:.4f}%）。',
              f'严格模型与LP目标差：{summary["lp_gap_yuan"]:.10f}元。',
              f'严格模型状态：{strict_info["message"]}；相对间隙{strict_info["relative_gap"]:.3g}；耗时{strict_info["elapsed_seconds"]:.4f}秒。','',
              '## 表1：指定时段及全天购电','', '| 时间段 | 购电量/kWh |','|---|---:|']
    for r in selected: lines.append(f'| {r["interval_label"]} | {r["purchase_kwh"]:.6f} |')
    lines += [f'| 全天 | {audits["strict"]["purchase_kwh"]:.6f} |',f'\n全天购电费：{audits["strict"]["cost_yuan"]:.6f}元。','',
              '## 表2：充放电量','', '| 时间段 | 充电量/kWh | 放电量/kWh |','|---|---:|---:|']
    for b in blocks: lines.append(f'| {b["interval_label"]} | {b["charge_kwh"]:.6f} | {b["discharge_kwh"]:.6f} |')
    lines += ['', '0:00与24:00储电量均为6000 kWh。','', '## 独立校验','',
              '- 逐段平衡、SOC递推/上下界、功率限制、充放电互斥、初末状态、全天能量恒等式及费用全部通过。',
              '- CSV与Excel写出后重新读取核对通过；注入1 kWh错误购电量能被校验器拒绝。',
              '- 原题及全部附件执行前后SHA256一致。',
              f'- 最大逐段能量平衡残差：{audits["strict"]["max_balance_error_kwh"]:.3g} kWh。',
              f'- SOC实际范围：{audits["strict"]["soc_min_kwh"]:.6f}—{audits["strict"]["soc_max_kwh"]:.6f} kWh。',
              f'- 充放电损耗：{audits["strict"]["storage_loss_kwh"]:.6f} kWh。','',
              '## 输出及解释边界','',
              '- results/result1.xlsx：保留两张要求的工作表，规范时间标签为0:00–24:00；原模板不修改。',
              '- results/q1/dispatch.csv：全精度逐段策略；four_hour_summary.csv与selected_intervals.csv为题目指定汇总。',
              '- results/q1/experiment.json：求解器状态、版本、目标界、审计、输入哈希。',
              '- 本结果对给定预测曲线最优，不代表存在预测误差时仍有相同费用或供电保障。',
              '- 未人为加入电池折旧或平滑成本；浮点容差内最优不意味着调度方案唯一。','']
    (ROOT/'reports/q1_experiment.md').write_text('\n'.join(lines),encoding='utf-8')
    print(json.dumps({k:v for k,v in summary.items() if k!='inputs'},ensure_ascii=False,indent=2))


if __name__=='__main__': main()
