"""Run Q2 v1 from January warm-up through February-December evaluation."""
from __future__ import annotations
import csv
import hashlib
import itertools
import json
import math
import sqlite3
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import openpyxl
import scipy
from openpyxl.utils import get_column_letter

from q2_model import Config, execute, idle_plan, plan, predict
from validate_q2 import audit

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results/q2'
EPS=1e-7


def write_csv(name, rows):
    with (OUT/f'{name}.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)


def label(start,end):
    return f'{start//60:02d}:{start%60:02d}-{end//60:02d}:{end%60:02d}'


def tune(actual,price):
    candidates=[]
    for method,pv_days,penalties in itertools.product(
        ('weekly_last','weekly_weighted'),(3,7),((.2,.05),(.5,.1),(1.,.2))):
        config=Config(method,pv_days,*penalties)
        soc=6000.;total=0.;emergency_total=0.
        for d in range(14,31):
            forecast=predict(actual[:d],config)
            p=plan(forecast,price,soc,config,final_soc=6000 if d==30 else None)
            e,w,cost=execute(p,actual[d],price)
            total+=float(cost.sum());emergency_total+=float(e.sum())
            soc+=float(.9*p['charge'].sum()-p['discharge'].sum()/.9)
        if abs(soc-6000)>1e-5:raise ValueError('Validation terminal SOC')
        candidates.append(dict(candidate_id=len(candidates),**config.record(),
                               validation_start='2025-01-15',validation_end='2025-01-31',
                               actual_cost_yuan=total,emergency_kwh=emergency_total))
    chosen=min(candidates,key=lambda r:(r['actual_cost_yuan'],r['candidate_id']))
    config=Config(**{k:chosen[k] for k in Config().record()})
    return config,candidates,chosen['candidate_id']


def execution_rows(ds,p,actual,price,soc,phase):
    midnight=datetime.fromisoformat(ds)
    emergency,waste,cost=execute(p,actual,price)
    rows=[]
    for t in range(144):
        end_soc=soc+.9*float(p['charge'][t])-float(p['discharge'][t])/.9
        row=dict(date=ds,slot=t+1,phase=phase,
                 interval_start=(midnight+timedelta(minutes=t*10)).isoformat(timespec='minutes'),
                 interval_end=(midnight+timedelta(minutes=(t+1)*10)).isoformat(timespec='minutes'),
                 interval_label=label(t*10,(t+1)*10),plan_issue_time=midnight.isoformat(timespec='minutes'),
                 history_available_through='' if ds=='2025-01-01' else midnight.isoformat(timespec='minutes'),
                 price_yuan_per_kwh=float(price[t]),load_forecast_kwh=float(p['forecast'][t,0]),
                 pv_forecast_kwh=float(p['forecast'][t,1]),load_kwh=float(actual[t,0]),pv_actual_kwh=float(actual[t,1]),
                 purchase_kwh=float(p['purchase'][t]),planned_charge_kwh=float(p['charge'][t]),
                 planned_discharge_kwh=float(p['discharge'][t]),charge_kwh=float(p['charge'][t]),
                 discharge_kwh=float(p['discharge'][t]),emergency_kwh=float(emergency[t]),
                 curtailment_kwh=float(waste[t]),soc_start_kwh=soc,soc_end_kwh=end_soc,
                 planned_cost_yuan=float(price[t]*p['purchase'][t]),
                 emergency_cost_yuan=float(5*price[t]*emergency[t]),total_cost_yuan=float(cost[t]))
        rows.append(row);soc=end_soc
    if np.max(np.abs(np.array([r['soc_start_kwh'] for r in rows]+[soc])-p['soc']))>1e-5:
        raise ValueError('Executed SOC differs from fixed plan')
    return rows,soc


def summarize(rows):
    def total(k):return math.fsum(r[k] for r in rows)
    return dict(days=len(rows)//144,purchase_kwh=total('purchase_kwh'),emergency_kwh=total('emergency_kwh'),
                actual_grid_kwh=total('purchase_kwh')+total('emergency_kwh'),
                planned_cost_yuan=total('planned_cost_yuan'),emergency_cost_yuan=total('emergency_cost_yuan'),
                total_cost_yuan=total('total_cost_yuan'),charge_kwh=total('charge_kwh'),
                discharge_kwh=total('discharge_kwh'),curtailment_kwh=total('curtailment_kwh'),
                emergency_intervals=sum(r['emergency_kwh']>EPS for r in rows),
                emergency_days=len({r['date'] for r in rows if r['emergency_kwh']>EPS}),
                emergency_while_charging_kwh=math.fsum(r['emergency_kwh'] for r in rows if r['charge_kwh']>EPS),
                emergency_while_charging_intervals=sum(r['emergency_kwh']>EPS and r['charge_kwh']>EPS for r in rows),
                waste_while_discharging_kwh=math.fsum(r['curtailment_kwh'] for r in rows if r['discharge_kwh']>EPS),
                soc_start_kwh=rows[0]['soc_start_kwh'],soc_end_kwh=rows[-1]['soc_end_kwh'])


def events(rows):
    result=[]
    for offset in range(0,len(rows),144):
        daily=rows[offset:offset+144];current=[]
        for r in daily+[None]:
            if r is not None and r['emergency_kwh']>EPS:
                current.append(r)
            elif current:
                result.append(dict(date=current[0]['date'],start_slot=current[0]['slot'],end_slot=current[-1]['slot'],
                                   interval_label=label((current[0]['slot']-1)*10,current[-1]['slot']*10),
                                   emergency_kwh=math.fsum(x['emergency_kwh'] for x in current)))
                current=[]
    if abs(math.fsum(r['emergency_kwh'] for r in result)-math.fsum(r['emergency_kwh'] for r in rows))>EPS*len(rows):
        raise ValueError('Event merge lost energy')
    return result


def workbook(rows,daily,event_rows,output_path=None):
    w=openpyxl.Workbook();p=w.active;p.title='计划购电量'
    p.append(['日期\\时间']+[label(t*10,(t+1)*10) for t in range(144)]+['全天购电量','全天购电费'])
    battery=w.create_sheet('充放电量');battery.append(['日期','时间段','充电量','放电量','时刻','储电量'])
    e=w.create_sheet('紧急购电量');e.append(['日期','购电时间段','购电量'])
    groups={d['date']:[] for d in daily}
    for r in event_rows:groups[r['date']].append(r)
    for i,d in enumerate(daily):
        rs=rows[i*144:(i+1)*144];date_value=datetime.fromisoformat(d['date'])
        p.append([date_value]+[r['purchase_kwh'] for r in rs]+[d['purchase_kwh'],d['planned_cost_yuan']])
        for j in range(6):
            block=rs[j*24:(j+1)*24]
            battery.append([date_value if j==0 else None,label(j*240,(j+1)*240),
                            math.fsum(r['charge_kwh'] for r in block),math.fsum(r['discharge_kwh'] for r in block),
                            '0:00' if j==0 else '24:00' if j==1 else None,
                            d['soc_start_kwh'] if j==0 else d['soc_end_kwh'] if j==1 else None])
        if not groups[d['date']]:e.append([date_value,None,0.])
        for j,r in enumerate(groups[d['date']]):
            e.append([date_value if j==0 else None,r['interval_label'],r['emergency_kwh']])
    for sheet in w:
        sheet.freeze_panes='B2'
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,datetime):cell.number_format='yyyy/mm/dd'
                elif isinstance(cell.value,(int,float)):cell.number_format='0.000000'
        for j in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(j)].width=19
    path=ROOT/'results/result2.xlsx' if output_path is None else Path(output_path)
    w.save(path);w.close()
    return path


def verify_export(path,rows,daily,event_rows):
    w=openpyxl.load_workbook(path,read_only=True,data_only=True)
    ps=list(w['计划购电量'].values)
    assert len(ps)==335 and len(ps[0])==147
    assert list(ps[0][1:145])==[label(t*10,(t+1)*10) for t in range(144)]
    for i,r in enumerate(ps[1:]):
        assert r[0].date().isoformat()==daily[i]['date']
        assert max(abs(a-b['purchase_kwh']) for a,b in zip(r[1:145],rows[i*144:(i+1)*144]))<1e-7
        assert abs(r[145]-daily[i]['purchase_kwh'])<1e-7 and abs(r[146]-daily[i]['planned_cost_yuan'])<1e-7
    bs=list(w['充放电量'].values)[1:]
    assert len(bs)==334*6
    for i,d in enumerate(daily):
        for j in range(6):
            r=bs[i*6+j];block=rows[i*144+j*24:i*144+(j+1)*24]
            assert r[1]==label(j*240,(j+1)*240)
            assert abs(r[2]-math.fsum(x['charge_kwh'] for x in block))<1e-7
            assert abs(r[3]-math.fsum(x['discharge_kwh'] for x in block))<1e-7
        assert abs(bs[i*6][5]-d['soc_start_kwh'])<1e-7 and abs(bs[i*6+1][5]-d['soc_end_kwh'])<1e-7
    es=list(w['紧急购电量'].values)[1:]
    written=[];date_value=None
    for r in es:
        if r[0] is not None:date_value=r[0].date().isoformat()
        if r[1] is not None:written.append((date_value,r[1],r[2]))
        else:assert r[2]==0
    assert len(written)==len(event_rows)
    for a,b in zip(written,event_rows):
        assert a[:2]==(b['date'],b['interval_label']) and abs(a[2]-b['emergency_kwh'])<1e-7
    w.close()


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    quality=json.loads((ROOT/'reports/data_quality_report.json').read_text(encoding='utf-8'))
    assert quality['status']=='PASS_WITH_WARNINGS'
    for r in quality['sources']:assert hashlib.sha256((ROOT/r['path']).read_bytes()).hexdigest()==r['sha256']
    db=ROOT/'data/processed/microgrid.sqlite'
    with sqlite3.connect(f'file:{db.as_posix()}?mode=ro',uri=True) as c:
        flat=c.execute('SELECT date,slot,load_kwh,pv_actual_kwh FROM actuals ORDER BY date,slot').fetchall()
        price=np.array([r[0] for r in c.execute('SELECT price_yuan_per_kwh FROM baseline ORDER BY slot')])
    assert len(flat)==365*144
    dates=[flat[i*144][0] for i in range(365)]
    actual=np.array([[r[2],r[3]] for r in flat]).reshape(365,144,2)
    # The selector only receives January: evaluation observations are inaccessible to it.
    selected,candidates,candidate_id=tune(actual[:31].copy(),price)
    print('January-only selection:',selected.record(),flush=True)
    write_csv('january_validation',candidates)
    all_rows=[];control=[];solver_rows=[];soc=6000.;default=Config()
    for d,ds in enumerate(dates):
        config=default if d<31 else selected
        if d==0:
            p=idle_plan(soc)
        else:
            forecast=predict(actual[:d].copy(),config)
            p=plan(forecast,price,soc,config,final_soc=6000 if d in (30,364) else None)
        # No target-day actual was passed to predict/plan. Only now execute against it.
        day_rows,soc=execution_rows(ds,p,actual[d],price,soc,'warmup' if d<31 else 'evaluation')
        all_rows.extend(day_rows)
        solver_rows.append(dict(date=ds,config=json.dumps(config.record(),sort_keys=True),
                                status=p['status'],relative_gap=p['gap'],seconds=p['seconds'],
                                predicted_objective_yuan=p['objective'],terminal_penalty_yuan=p['terminal_penalty'],
                                hard_terminal_soc=int(d in (30,364))))
        if d>=31:
            idle=idle_plan(6000.,p['forecast'])
            cr,_=execution_rows(ds,idle,actual[d],price,6000.,'evaluation');control.extend(cr)
        if (d+1)%60==0:print(f'Executed through {ds}',flush=True)
    formal=all_rows[31*144:];warm=all_rows[:31*144]
    validation=dict(full=audit(all_rows),evaluation=audit(formal),no_storage=audit(control))
    # Sentinel defect proves the audit is not merely checking solver status.
    broken=[dict(r) for r in formal[:144]];broken[0]['purchase_kwh']+=1
    try:audit(broken,require_boundaries=False)
    except ValueError:validation['injected_imbalance_rejected']=True
    else:raise ValueError('Auditor accepted injected imbalance')
    # Exercise cold-start and short-history fallbacks without using future observations.
    assert np.array_equal(predict(actual[:1],selected)[:,0],actual[0,:,0])
    prefix=actual[:31].copy();forecast_before=predict(prefix,selected);altered=actual.copy();altered[31:]=1e9
    assert np.array_equal(forecast_before,predict(altered[:31],selected))
    validation['future_suffix_invariance']=True
    daily=[]
    for i in range(334):
        rs=formal[i*144:(i+1)*144];base=control[i*144:(i+1)*144]
        daily.append(dict(date=rs[0]['date'],**summarize(rs),no_storage_total_cost_yuan=summarize(base)['total_cost_yuan']))
    ev=events(formal)
    write_csv('execution',all_rows);write_csv('daily_summary',daily);write_csv('solver_log',solver_rows)
    write_csv('no_storage_daily',[dict(date=control[i*144]['date'],**summarize(control[i*144:(i+1)*144])) for i in range(334)])
    write_csv('emergency_events',ev if ev else [dict(date='',start_slot='',end_slot='',interval_label='',emergency_kwh=0.)])
    # Revalidate persisted execution, not just in-memory arrays.
    string_fields={'date','phase','interval_start','interval_end','interval_label','plan_issue_time','history_available_through'}
    with (OUT/'execution.csv').open(encoding='utf-8-sig',newline='') as f:
        persisted=[{k:v if k in string_fields else int(v) if k=='slot' else float(v) for k,v in r.items()} for r in csv.DictReader(f)]
    validation['persisted_execution']=audit(persisted)
    path=workbook(formal,daily,ev);verify_export(path,formal,daily,ev)
    validation['workbook_readback']=True
    for r in quality['sources']:assert hashlib.sha256((ROOT/r['path']).read_bytes()).hexdigest()==r['sha256']
    validation['originals_unchanged']=True
    main_summary=summarize(formal);base_summary=summarize(control)
    savings=base_summary['total_cost_yuan']-main_summary['total_cost_yuan']
    forecast_metrics={
        'load_mae_kw':math.fsum(abs(r['load_forecast_kwh']-r['load_kwh']) for r in formal)/len(formal)*6,
        'pv_mae_kw':math.fsum(abs(r['pv_forecast_kwh']-r['pv_actual_kwh']) for r in formal)/len(formal)*6,
        'net_mae_kw':math.fsum(abs(r['load_forecast_kwh']-r['pv_forecast_kwh']-r['load_kwh']+r['pv_actual_kwh']) for r in formal)/len(formal)*6,
    }
    result=dict(version='q2-v1-point-fixed-plan',selected_config=selected.record(),selected_candidate_id=candidate_id,
                evaluation_start='2025-02-01',evaluation_end='2025-12-31',formal=main_summary,
                no_storage=base_summary,warmup=summarize(warm),savings_yuan=savings,
                savings_percent=savings/base_summary['total_cost_yuan']*100,forecast_metrics=forecast_metrics,
                validation=validation,emergency_event_count=len(ev),worst_day=max(daily,key=lambda r:r['total_cost_yuan']),
                solver_max_relative_gap=max(r['relative_gap'] for r in solver_rows),
                solver_total_seconds=math.fsum(r['seconds'] for r in solver_rows),scipy_version=scipy.__version__,
                source_hashes=quality['sources'],database_sha256=hashlib.sha256(db.read_bytes()).hexdigest())
    (OUT/'experiment.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    lines=['# 第二问v1实验报告','',
           '这是点预测加固定日前储能计划的因果基线；未实现误差场景或日内储能重调度。',
           '每日确定性MILP求解最优，不代表全年不确定性问题的最优策略。','',
           '## 信息与实验边界','',
           '- 仅附件1固定电价和附件2历史实际值。预测/计划完成后才揭示目标日实际数据；不使用附件3/4。',
           '- 正式结果为2月1日—12月31日334天；1月按事先固定参数预热，费用单列。',
           '- 1月15—31日顺序验证12组参数，候选统一6000 kWh起止；2月1日锁定，后续仅增加历史、不重选参数。',
           '- 正式运行1月31日及12月31日以实际调度回到6000，保证对照边界一致；不是每天强制循环或重置。',
           '- 计划购电费用不退；紧急购电额外按5倍电价结算。SOC软惩罚不计入实际费用。',
           f'- 选中参数：`{json.dumps(selected.record(),ensure_ascii=False)}`。','',
           '## 正式期对照','', '| 指标 | 点预测储能策略 | 同预测无储能 |','|---|---:|---:|']
    for key,title in [('purchase_kwh','计划购电/kWh'),('emergency_kwh','紧急购电/kWh'),('planned_cost_yuan','计划费/元'),
                      ('emergency_cost_yuan','紧急费/元'),('total_cost_yuan','总费用/元'),('curtailment_kwh','富余弃用/kWh')]:
        lines.append(f'| {title} | {main_summary[key]:.6f} | {base_summary[key]:.6f} |')
    lines += ['',f'相对同预测无储能节费{savings:.6f}元（{result["savings_percent"]:.4f}%）。',
              f'正式期紧急购电{main_summary["emergency_days"]}天、{main_summary["emergency_intervals"]}个区间，共{len(ev)}段连续事件。',
              f'1月预热费用：{result["warmup"]["total_cost_yuan"]:.6f}元，不计入上述正式期。',
              f'预测误差MAE：负载{forecast_metrics["load_mae_kw"]:.4f} kW，光伏{forecast_metrics["pv_mae_kw"]:.4f} kW，净负载{forecast_metrics["net_mae_kw"]:.4f} kW。','',
              '## 固定计划的局限','',
              f'- 充电同时发生紧急购电：{main_summary["emergency_while_charging_intervals"]}个区间；这些区间紧急电量共{main_summary["emergency_while_charging_kwh"]:.6f} kWh（不代表全部用于充电）。',
              f'- 放电区间的富余弃用共{main_summary["waste_while_discharging_kwh"]:.6f} kWh（不代表全部来自放电）。',
              '- 固定计划不根据区间结束真实值修改已发出的控制，因此出现上述现象；后续可评估明确因果信息下的日内重调度及区间内保护。',
              '- 1月验证样本有限；全年季节变化可能降低预测与参数泛化质量。',
              '- 沿用第一问右端代表功率、充放电各90%、设备外部电量计量的假设。','',
              '## 指定日期','']
    for ds in ('2025-03-20','2025-06-21','2025-09-23','2025-12-21'):
        rs=[r for r in formal if r['date']==ds];day_summary=next(r for r in daily if r['date']==ds)
        lines += [f'### {ds}','', '| 时间段 | 计划购电/kWh |','|---|---:|']
        for h in (10,12,14,16,18,20):lines.append(f'| {rs[h*6]["interval_label"]} | {rs[h*6]["purchase_kwh"]:.6f} |')
        lines += ['',f'全天计划购电{day_summary["purchase_kwh"]:.6f} kWh；计划费{day_summary["planned_cost_yuan"]:.6f}元；紧急购电{day_summary["emergency_kwh"]:.6f} kWh；紧急费{day_summary["emergency_cost_yuan"]:.6f}元；总费用{day_summary["total_cost_yuan"]:.6f}元。','',
                  '| 时间段 | 充电/kWh | 放电/kWh |','|---|---:|---:|']
        for j in range(6):
            block=rs[j*24:(j+1)*24]
            lines.append(f'| {label(j*240,(j+1)*240)} | {math.fsum(r["charge_kwh"] for r in block):.6f} | {math.fsum(r["discharge_kwh"] for r in block):.6f} |')
        lines += ['',f'0:00储电量{day_summary["soc_start_kwh"]:.6f} kWh；24:00储电量{day_summary["soc_end_kwh"]:.6f} kWh。','', '| 紧急购电时间段 | 电量/kWh |','|---|---:|']
        de=[e for e in ev if e['date']==ds]
        lines += [f'| {e["interval_label"]} | {e["emergency_kwh"]:.6f} |' for e in de] or ['| 无 | 0 |']
        lines.append('')
    lines += ['## 验证与复现','',
              '- 全年逐段供需、SOC递推、边界、功率、互斥、费用、时间连续与历史截止检查通过。',
              '- 写出CSV重新物理审计；Excel逐日计划、6组充放电、SOC及事件回读通过；紧急事件合并电量守恒。',
              '- 修改未来数据不改变固定历史前缀的预测；注入1kWh购电错误被审计拒绝；原件哈希保持不变。',
              f'- 最大MILP相对间隙{result["solver_max_relative_gap"]:.3g}；全年求解器累计耗时{result["solver_total_seconds"]:.3f}秒（不含参数验证/导出）。',
              '- 执行 `python scripts/solve_q2.py`；结果 `results/result2.xlsx`，完整明细及参数选择见 `results/q2/`。',
              '- 工作簿“全天购电量/费”为计划量/计划费；实际总量、紧急费、总费用在daily_summary.csv独立列示。','']
    (ROOT/'reports/q2_experiment.md').write_text('\n'.join(lines),encoding='utf-8')
    print(json.dumps({k:result[k] for k in ('selected_config','formal','savings_yuan','savings_percent','forecast_metrics','solver_total_seconds')},ensure_ascii=False,indent=2))


if __name__=='__main__':main()
