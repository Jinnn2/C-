"""Reproducible local normalization; originals are read-only inputs."""
from __future__ import annotations

import csv
import hashlib
import json
import math
import sqlite3
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'data' / 'processed'
REPORTS = ROOT / 'reports'
CHECKS: list[dict] = []
WARNINGS: list[str] = []
DAYS = [(date(2025, 1, 1) + timedelta(days=i)).isoformat() for i in range(365)]


def check(name, condition, detail=''):
    CHECKS.append({'check': name, 'passed': bool(condition), 'detail': str(detail)})
    if not condition:
        raise ValueError(f'{name}: {detail}')


def iso(value):
    return value.isoformat(timespec='minutes')


def minute(value):
    if isinstance(value, time):
        check('时间秒数为0', value.second == 0 and value.microsecond == 0, value)
        return value.hour * 60 + value.minute
    value = str(value).strip()
    extra = 1440 if value.endswith('+1') else 0
    value = value.removesuffix('+1')
    h, m = map(int, value.split(':'))
    check('时间格式范围', 0 <= h <= 24 and 0 <= m < 60 and (h != 24 or m == 0), value)
    return extra + h * 60 + m


def day(value):
    if isinstance(value, datetime):
        check('日期无日内偏移', value.time() == time(0), value)
        return value.date().isoformat()
    y, m, d = map(int, str(value).strip().replace('/', '-').split('-'))
    return date(y, m, d).isoformat()


def number(value, location):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f'非数值/缺失: {location}: {value!r}')
    value = float(value)
    if not math.isfinite(value) or value < 0:
        raise ValueError(f'非有限值/负值: {location}: {value}')
    return value


def source(file, sheet, row, col):
    return f'附件/{file}::{sheet}!{get_column_letter(col)}{row}'


def read(file, sheet):
    w = openpyxl.load_workbook(ROOT / '附件' / file, read_only=True, data_only=True)
    rows = list(w[sheet].values)
    w.close()
    return rows


def manifest():
    files = [ROOT / 'C题.pdf', *sorted((ROOT / '附件').rglob('*.xlsx'))]
    return [dict(path=p.relative_to(ROOT).as_posix(), bytes=p.stat().st_size,
                 sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in files]


def matrix(file, sheet):
    rows = read(file, sheet)
    check(f'{file}/{sheet}尺寸', len(rows) == 366 and all(len(r) == 145 for r in rows))
    check(f'{file}/{sheet}时间轴', [minute(t) for t in rows[0][1:]] == list(range(10, 1441, 10)))
    check(f'{file}/{sheet}日期连续唯一', [day(r[0]) for r in rows[1:]] == DAYS)
    return [[number(v, source(file, sheet, i + 2, j + 2)) for j, v in enumerate(r[1:])]
            for i, r in enumerate(rows[1:])]


def label(slot):
    def fmt(m):
        return f'{m // 60:02d}:{m % 60:02d}'
    return f'{fmt((slot - 1) * 10)}-{fmt(slot * 10)}'


def export_csv(name, rows):
    check(f'{name}非空', bool(rows))
    path = OUT / f'{name}.csv'
    temp = path.with_suffix('.csv.tmp')
    with temp.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    temp.replace(path)


def stats(values):
    values = sorted(values)
    n = len(values)
    def quantile(p):
        x = (n - 1) * p
        low = int(x)
        return values[low] + (x - low) * (values[min(low + 1, n - 1)] - values[low])
    return dict(count=n, min=values[0], p01=quantile(.01), median=quantile(.5),
                p99=quantile(.99), max=values[-1], mean=math.fsum(values) / n,
                zeros=values.count(0.0))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    REPORTS.mkdir(parents=True, exist_ok=True)
    before = manifest()
    # The workbook values are authoritative; formulas without cached numbers fail.
    rows = read('附件1.xlsx', 'Sheet1')
    check('附件1尺寸', len(rows) == 145 and all(len(r) == 4 for r in rows))
    check('附件1时间轴', [minute(r[0]) for r in rows[1:]] == list(range(10, 1441, 10)))
    baseline = []
    for k, row in enumerate(rows[1:], 1):
        p, l, g = [number(v, source('附件1.xlsx', 'Sheet1', k + 1, j))
                   for j, v in enumerate(row[1:], 2)]
        baseline.append(dict(slot=k, start_minute=(k-1)*10, end_minute=k*10,
                             interval_label=label(k), price_yuan_per_kwh=p,
                             load_kw=l, pv_forecast_kw=g, load_kwh=l/6,
                             pv_forecast_kwh=g/6, price_source=source('附件1.xlsx', 'Sheet1', k+1, 2),
                             load_source=source('附件1.xlsx', 'Sheet1', k+1, 3),
                             pv_source=source('附件1.xlsx', 'Sheet1', k+1, 4)))
    loads = matrix('附件2.xlsx', '小区负载')
    pvs = matrix('附件2.xlsx', '光伏发电实际功率')
    prices = matrix('附件4.xlsx', 'Sheet1')
    actuals, daily = [], []
    for i, ds in enumerate(DAYS):
        midnight = datetime.fromisoformat(ds)
        for j in range(144):
            start, end = midnight + timedelta(minutes=j*10), midnight + timedelta(minutes=(j+1)*10)
            actuals.append(dict(date=ds, slot=j+1, interval_start=iso(start), interval_end=iso(end),
                                observed_available_at=iso(end), price_available_at=iso(end),
                                load_kw=loads[i][j], pv_actual_kw=pvs[i][j],
                                load_kwh=loads[i][j]/6, pv_actual_kwh=pvs[i][j]/6,
                                net_load_kwh=(loads[i][j]-pvs[i][j])/6,
                                price_actual_yuan_per_kwh=prices[i][j],
                                price_fixed_yuan_per_kwh=baseline[j]['price_yuan_per_kwh'],
                                load_source=source('附件2.xlsx', '小区负载', i+2, j+2),
                                pv_source=source('附件2.xlsx', '光伏发电实际功率', i+2, j+2),
                                price_source=source('附件4.xlsx', 'Sheet1', i+2, j+2)))
        daily.append(dict(date=ds, load_kwh=math.fsum(loads[i])/6, pv_actual_kwh=math.fsum(pvs[i])/6,
                          net_load_kwh=math.fsum(l-g for l,g in zip(loads[i], pvs[i]))/6,
                          pv_surplus_intervals=sum(g>l for l,g in zip(loads[i], pvs[i])),
                          price_min=min(prices[i]), price_max=max(prices[i]),
                          price_mean=math.fsum(prices[i])/144))
    check('实际区间数', len(actuals) == 52560)
    check('实际区间连续', all(a['interval_end'] == b['interval_start'] for a,b in zip(actuals, actuals[1:])))
    check('实际区间唯一', len({r['interval_end'] for r in actuals}) == len(actuals))
    check('功率电量换算', all(abs(r['load_kwh']*6-r['load_kw'])<1e-9 and
                                 abs(r['pv_actual_kwh']*6-r['pv_actual_kw'])<1e-9 for r in actuals))
    observed = {r['interval_end']: r['pv_actual_kw'] for r in actuals}
    rows = read('附件3.xlsx', 'Sheet1')
    check('预报尺寸', len(rows) == 1461 and all(len(r) == 26 for r in rows))
    check('预报提前量表头', list(rows[0][2:]) == [f'预报{i}小时' for i in range(1,25)])
    hourly, fine, issues = [], [], []
    last_day = None
    for rownum, row in enumerate(rows[1:], 2):
        if row[0] not in ('', None):
            last_day = day(row[0])
        check('预报日期可解析', last_day is not None, rownum)
        m = minute(row[1])
        check('预报发布时刻', m in (0, 360, 720, 1080), rownum)
        issue = datetime.fromisoformat(last_day) + timedelta(minutes=m)
        issues.append(iso(issue))
        vals = [number(v, source('附件3.xlsx', 'Sheet1', rownum, j+3)) for j,v in enumerate(row[2:])]
        for h, v in enumerate(vals, 1):
            target = iso(issue + timedelta(hours=h))
            hourly.append(dict(issue_time=iso(issue), lead_hours=h, target_time=target,
                               pv_forecast_kw=v, actual_target_available=int(target in observed),
                               source=source('附件3.xlsx', 'Sheet1', rownum, h+2)))
        for k in range(1,145):
            # No actual observations are read here: all inputs belong to this issuance.
            if k <= 6:
                left = right = 1
                v = vals[0]
                method = 'first_hour_hold' if k < 6 else 'hourly_anchor'
            elif k % 6 == 0:
                left = right = k // 6
                v = vals[left-1]
                method = 'hourly_anchor'
            else:
                left, right = k // 6, k // 6 + 1
                fraction = (k % 6) / 6
                v = vals[left-1] * (1-fraction) + vals[right-1] * fraction
                method = 'linear_between_hours'
            end = iso(issue + timedelta(minutes=k*10))
            fine.append(dict(issue_time=iso(issue), lead_minutes=k*10,
                             interval_start=iso(issue+timedelta(minutes=(k-1)*10)), interval_end=end,
                             pv_forecast_kw=v, pv_forecast_kwh=v/6, method=method,
                             actual_target_available=int(end in observed),
                             left_source=source('附件3.xlsx', 'Sheet1', rownum, left+2),
                             right_source=source('附件3.xlsx', 'Sheet1', rownum, right+2)))
    expected = [iso(datetime.fromisoformat(d)+timedelta(hours=h)) for d in DAYS for h in (0,6,12,18)]
    check('全年每日4次发布且无重复', issues == expected)
    check('原始预报总数', len(hourly) == 35040)
    check('派生预报总数', len(fine) == 210240)
    check('预报目标均在发布后', all(r['issue_time'] < r['interval_end'] for r in fine))
    anchors = {(r['issue_time'],r['target_time']):r['pv_forecast_kw'] for r in hourly}
    check('插值严格保留全部原始整点', all(r['pv_forecast_kw'] == anchors[(r['issue_time'],r['interval_end'])]
                                                     for r in fine if r['lead_minutes'] % 60 == 0))
    check('派生预报非负有限', all(math.isfinite(r['pv_forecast_kw']) and r['pv_forecast_kw'] >= 0 for r in fine))
    beyond = sum(not r['actual_target_available'] for r in hourly)
    WARNINGS.append(f'原始小时预报有{beyond}条目标超出实际观测范围，均保留，不补造2026年实际值。')
    mappings = []
    for path in sorted((ROOT/'附件'/'附件5').glob('*.xlsx')):
        w = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for s in w:
            rows = list(s.values)
            if s.title not in ('计划购电量', '调整购电量'):
                continue
            if path.name == 'result1.xlsx':
                check('result1时段数', len(rows)-1 == 144)
                labels = [(i+2, 1, r[0]) for i,r in enumerate(rows[1:])]
            else:
                check(f'{path.name}/{s.title}输出日期', [day(r[0]) for r in rows[1:]] == DAYS[31:])
                check(f'{path.name}/{s.title}列数', len(rows[0]) == 147)
                labels = [(1, j+2, v) for j,v in enumerate(rows[0][1:145])]
            for k,(rownum,colnum,original) in enumerate(labels,1):
                text = str(original)
                try:
                    a,b = text.split('-')
                    matches = minute(a)==(k-1)*10 and minute(b)==k*10
                except (ValueError, TypeError):
                    matches = False
                mappings.append(dict(file=path.name, sheet=s.title,
                                     cell=f'{get_column_letter(colnum)}{rownum}', slot=k,
                                     original_label=text, normalized_label=label(k),
                                     correction_required=int(not matches), status='proposed_order_mapping'))
        w.close()
    mismatch = sum(r['correction_required'] for r in mappings)
    WARNINGS.append(f'模板时间映射共{len(mappings)}条，{mismatch}条与规范顺序时段不一致；保留原件，仅提供建议映射。')
    WARNINGS.append('模板充放电/紧急购电工作表是示例布局，不覆盖全部输出日，后续必须动态扩展。')
    WARNINGS.append('原始时间标注混合Excel时间对象与字符串，已解析为统一时间；右端区间代表值属于建模假设。')
    WARNINGS.append('10分钟预报首小时保持第1小时值，之后线性插值；派生预报不是额外观测，需后续检验插值敏感性。')
    WARNINGS.append('未来实际电价及负载/光伏不得作为日前已知值；1月预热SOC和调整结算口径待求解阶段固定。')
    metadata = [dict(key=k,value=json.dumps(v,ensure_ascii=False)) for k,v in {
        'timezone':'Asia/Shanghai', 'interval_minutes':10, 'power_unit':'kW', 'energy_unit':'kWh',
        'price_unit':'yuan/kWh', 'interval_semantics':'right_endpoint_representative_assumption',
        'forecast_first_hour':'hold_first_available_forecast_no_actual_data',
        'forecast_later_hours':'linear_endpoint_interpolation',
        'initial_soc_at_2025_01_01_kwh':6000, 'soc_min_kwh':1200, 'soc_max_kwh':10800,
        'nameplate_capacity_kwh':12000, 'max_grid_side_power_kw':5000,
        'charge_efficiency':.9, 'discharge_efficiency':.9,
        'output_start_date':'2025-02-01', 'output_end_date':'2025-12-31',
        'actual_price_availability':'assumed_interval_end_not_source_release_metadata',
        'template_mapping_status':'proposed_order_correction_not_official_clarification',
    }.items()]
    tables = dict(baseline=baseline, actuals=actuals, forecast_hourly=hourly,
                  forecast_10min=fine, template_time_mapping=mappings, metadata=metadata,
                  source_files=before)
    temp_db = OUT/'microgrid.build.sqlite'
    conn = sqlite3.connect(temp_db)
    keys = {'baseline':('slot',), 'actuals':('date','slot'),
            'forecast_hourly':('issue_time','lead_hours'), 'forecast_10min':('issue_time','lead_minutes'),
            'template_time_mapping':('file','sheet','cell'), 'metadata':('key',), 'source_files':('path',)}
    try:
        for name, data in tables.items():
            conn.execute(f'DROP TABLE IF EXISTS "{name}"')
            columns = list(data[0])
            types = {c:('INTEGER' if isinstance(data[0][c],int) else 'REAL' if isinstance(data[0][c],float) else 'TEXT') for c in columns}
            defs = ','.join(f'"{c}" {types[c]} NOT NULL' for c in columns)
            pk = ','.join(f'"{c}"' for c in keys[name])
            conn.execute(f'CREATE TABLE "{name}" ({defs}, PRIMARY KEY ({pk}))')
            placeholders = ','.join('?' for _ in columns)
            conn.executemany(f'INSERT INTO "{name}" VALUES ({placeholders})', ([r[c] for c in columns] for r in data))
            check(f'SQLite/{name}行数', conn.execute(f'SELECT COUNT(*) FROM "{name}"').fetchone()[0] == len(data))
        conn.execute('CREATE INDEX IF NOT EXISTS actual_end ON actuals(interval_end)')
        conn.execute('CREATE INDEX IF NOT EXISTS forecast_target ON forecast_hourly(target_time,issue_time)')
        conn.execute('CREATE INDEX IF NOT EXISTS forecast_fine_target ON forecast_10min(interval_end,issue_time)')
        check('SQLite完整性', conn.execute('PRAGMA integrity_check').fetchone()[0] == 'ok')
        # A known boundary lookup independently verifies slot/date ownership and conversion.
        first = conn.execute('SELECT interval_start,interval_end,load_kw,load_kwh FROM actuals WHERE date=? AND slot=1', (DAYS[0],)).fetchone()
        check('首段边界及实值回查', first[:2] == ('2025-01-01T00:00','2025-01-01T00:10') and
              abs(first[2]-3529.7296)<1e-9 and abs(first[3]-3529.7296/6)<1e-9)
        last = conn.execute('SELECT interval_start,interval_end FROM actuals WHERE date=? AND slot=144',(DAYS[-1],)).fetchone()
        check('年末归属回查', last == ('2025-12-31T23:50','2026-01-01T00:00'))
        conn.commit()
    finally:
        conn.close()
    check('原题及附件SHA256未变化', before == manifest())
    temp_db.replace(OUT/'microgrid.sqlite')
    for name,data in tables.items():
        export_csv(name,data)
    export_csv('daily_summary',daily)
    summary = {
        'baseline_price_yuan_per_kwh':stats([r['price_yuan_per_kwh'] for r in baseline]),
        'baseline_load_kw':stats([r['load_kw'] for r in baseline]),
        'baseline_pv_forecast_kw':stats([r['pv_forecast_kw'] for r in baseline]),
        'load_kw':stats([r['load_kw'] for r in actuals]),
        'pv_actual_kw':stats([r['pv_actual_kw'] for r in actuals]),
        'price_actual_yuan_per_kwh':stats([r['price_actual_yuan_per_kwh'] for r in actuals]),
        'pv_hourly_forecast_kw':stats([r['pv_forecast_kw'] for r in hourly]),
    }
    energy = dict(annual_load_kwh=math.fsum(r['load_kwh'] for r in actuals),
                  annual_pv_kwh=math.fsum(r['pv_actual_kwh'] for r in actuals),
                  pv_surplus_intervals=sum(r['net_load_kwh']<0 for r in actuals),
                  baseline_load_kwh=math.fsum(r['load_kwh'] for r in baseline),
                  baseline_pv_kwh=math.fsum(r['pv_forecast_kwh'] for r in baseline))
    check('逐日全年能量汇总一致', abs(energy['annual_load_kwh']-math.fsum(r['load_kwh'] for r in daily))<1e-6 and
          abs(energy['annual_pv_kwh']-math.fsum(r['pv_actual_kwh'] for r in daily))<1e-6)
    # Diagnostics are descriptive only; they are not outlier corrections or validation scores.
    differences = [abs(b['price_actual_yuan_per_kwh']-a['price_actual_yuan_per_kwh']) for a,b in zip(actuals,actuals[1:])]
    report = dict(status='PASS_WITH_WARNINGS', check_count=len(CHECKS), checks=CHECKS,
                  warnings=WARNINGS, statistics=summary, energy=energy,
                  hourly_forecasts_without_actual=beyond,
                  fine_forecasts_without_actual=sum(not r['actual_target_available'] for r in fine),
                  max_adjacent_price_change=max(differences), sources=before,
                  table_rows={name:len(data) for name,data in tables.items()})
    (REPORTS/'data_quality_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    lines = ['# 数据质量检查报告','', '状态：**通过，含需说明的警告**。未求解购电策略。', '',
             f'共执行{len(CHECKS)}项结构、时间、数值、换算、数据库和原件校验；全部通过。', '',
             '## 覆盖与可追溯性','',
             '- 2025年365天，每天144段；全年52560段，时间连续、日期与时段无重复。',
             '- 输出期2025-02-01至2025-12-31，共334天；1月保留为历史与预热数据。',
             '- 原始小时预报35040条；派生10分钟预报210240条，所有整点保持原值。',
             '- 数值区无缺失、非数值、非有限值或负数；不自动删除统计极值。',
             '- SQLite主键与完整性检查通过，全部原题及Excel的SHA256执行前后相同。', '',
             '## 分布统计','', '| 变量 | 数量 | 最小值 | 中位数 | 99%分位 | 最大值 |',
             '|---|---:|---:|---:|---:|---:|']
    for name,s in summary.items():
        lines.append(f"| {name} | {s['count']} | {s['min']:.4f} | {s['median']:.4f} | {s['p99']:.4f} | {s['max']:.4f} |")
    lines += ['', '## 能量检查','',
              f"- 全年负载：{energy['annual_load_kwh']:.4f} kWh。",
              f"- 全年光伏：{energy['annual_pv_kwh']:.4f} kWh。",
              f"- 光伏大于负载：{energy['pv_surplus_intervals']}个10分钟段；这是合理富余，不是负值错误。",
              f"- 附件1全天负载：{energy['baseline_load_kwh']:.4f} kWh；光伏预测：{energy['baseline_pv_kwh']:.4f} kWh。",
              '- 功率/6与电量逐条一致；逐日汇总与全年汇总一致。',
              f'- 相邻时段最大电价变化：{max(differences):.4f} 元/kWh；仅描述，不作为删除依据。', '',
              '## 警告与后续处理', ''] + [f'- {v}' for v in WARNINGS]
    lines += ['', '## 本阶段采用的假设','',
              '- 右端点功率代表此前10分钟区间；原题未明确平均/瞬时性质，保留为假设。',
              '- 实际值按区间结束可用；未来实际电价不提前公布。',
              '- 充放电各90%，电网侧限功率5000 kW；储电量1200—10800 kWh。',
              '- 模板仅提供规范顺序映射，未修改原标签或填写任何优化结果。', '',
              '## 使用方式','',
              '- 运行 `python scripts/prepare_data.py` 重建。',
              '- 数据库：`data/processed/microgrid.sqlite`；CSV在同目录。',
              '- 详细检查、文件哈希、统计：`reports/data_quality_report.json`。',
              '- 契约与字段含义：`docs/数据口径.md`；建模思路：`docs/建模方案.md`。', '']
    (REPORTS/'data_quality_report.md').write_text('\n'.join(lines),encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k in ('status','check_count','table_rows','energy','hourly_forecasts_without_actual')},ensure_ascii=False,indent=2))


if __name__ == '__main__':
    try:
        main()
    except Exception as exc:
        REPORTS.mkdir(parents=True,exist_ok=True)
        (REPORTS/'data_quality_report.json').write_text(json.dumps(dict(status='FAIL',error=str(exc),checks=CHECKS),ensure_ascii=False,indent=2),encoding='utf-8')
        (REPORTS/'data_quality_report.md').write_text(f'# 数据检查失败\n\n{exc}\n\n请勿使用此前的派生数据。\n',encoding='utf-8')
        raise
